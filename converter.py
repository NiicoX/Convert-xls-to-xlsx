"""Conversión de libros Excel .xls (BIFF) a .xlsx."""

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Tuple

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, Protection
from openpyxl.utils import get_column_letter
from xlrd.sheet import Sheet

XLS_EXTENSIONS = {".xls"}


@dataclass
class ConvertResult:
    source: Path
    destination: Path
    ok: bool
    engine: str
    message: str = ""


def is_xls_file(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in XLS_EXTENSIONS


def discover_xls(root: Path, recursive: bool) -> List[Path]:
    if root.is_file():
        return [root] if is_xls_file(root) else []
    pattern = "**/*.xls" if recursive else "*.xls"
    files = [p for p in root.glob(pattern) if p.is_file() and p.suffix.lower() == ".xls"]
    return sorted(files)


def destination_for(source: Path, output_dir: Optional[Path] = None) -> Path:
    folder = output_dir if output_dir is not None else source.parent
    return folder / (source.stem + ".xlsx")


def convert_file(
    source: Path,
    destination: Optional[Path] = None,
    overwrite: bool = False,
) -> ConvertResult:
    source = source.expanduser().resolve()
    dest = (destination or destination_for(source)).expanduser().resolve()
    dest.parent.mkdir(parents=True, exist_ok=True)

    if not is_xls_file(source):
        return ConvertResult(source, dest, False, "-", "No es un archivo .xls")

    if dest.exists() and not overwrite:
        return ConvertResult(source, dest, False, "-", "Ya existe el .xlsx (usa sobrescribir)")

    # Intentamos primero directamente con xlrd ya que este tipo de .xls es nativo de xlrd
    try:
        _convert_with_xlrd(source, dest)
        return ConvertResult(source, dest, True, "xlrd", "Convertido con xlrd/openpyxl")
    except Exception as exc:
        xlrd_error = str(exc)

    # Si por algo fallara xlrd, intentamos con Excel COM como respaldo
    try:
        _convert_with_excel(source, dest)
        return ConvertResult(source, dest, True, "Excel COM", "Convertido con Microsoft Excel")
    except Exception as exc:
        excel_error = str(exc)
        detail = "xlrd: {}; Excel: {}".format(_short(xlrd_error), _short(excel_error))
        return ConvertResult(source, dest, False, "-", detail)


def convert_many(
    sources: List[Path],
    output_dir: Optional[Path] = None,
    overwrite: bool = False,
) -> List[ConvertResult]:
    results = []
    for source in sources:
        dest = destination_for(source, output_dir)
        results.append(convert_file(source, dest, overwrite=overwrite))
    return results


def _short(text: str, limit: int = 160) -> str:
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    return text[: limit - 3] + "..."


def _convert_with_excel(source: Path, dest: Path) -> None:
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("pywin32 no está instalado") from exc

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        workbook = excel.Workbooks.Open(
            str(source),
            UpdateLinks=0,
            ReadOnly=True,
            IgnoreReadOnlyRecommended=True,
        )
        if dest.exists():
            os.remove(dest)
        workbook.SaveAs(str(dest), FileFormat=51)
        workbook.Close(SaveChanges=False)
        workbook = None
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def _convert_with_xlrd(source: Path, dest: Path) -> None:
    try:
        book = xlrd.open_workbook(str(source), formatting_info=True)
    except (xlrd.XLRDError, NotImplementedError):
        book = xlrd.open_workbook(str(source), formatting_info=False)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    xf_list = getattr(book, "xf_list", [])
    font_list = getattr(book, "font_list", [])
    format_map = getattr(book, "format_map", {})

    for sheet in book.sheets():
        worksheet = workbook.create_sheet(title=_safe_sheet_name(sheet.name))
        _copy_sheet(sheet, worksheet, book.datemode, xf_list, font_list, format_map)

    if dest.exists():
        os.remove(dest)
    workbook.save(str(dest))


def _safe_sheet_name(name: str) -> str:
    invalid = set(r"[]:*?/\\")
    cleaned = "".join("_" if ch in invalid else ch for ch in name).strip() or "Hoja"
    return cleaned[:31]


def _copy_sheet(sheet: Sheet, worksheet, datemode, xf_list, font_list, format_map) -> None:
    colinfo_map = getattr(sheet, "colinfo_map", {}) or {}
    for col_index in range(sheet.ncols):
        info = colinfo_map.get(col_index)
        if info is not None and getattr(info, "width", None):
            worksheet.column_dimensions[get_column_letter(col_index + 1)].width = min(
                info.width / 256.0, 80
            )

    merged = set()
    for rlo, rhi, clo, chi in sheet.merged_cells:
        worksheet.merge_cells(
            start_row=rlo + 1,
            start_column=clo + 1,
            end_row=rhi,
            end_column=chi,
        )
        for row in range(rlo, rhi):
            for col in range(clo, chi):
                if row == rlo and col == clo:
                    continue
                merged.add((row, col))

    for row_index in range(sheet.nrows):
        for col_index in range(sheet.ncols):
            if (row_index, col_index) in merged:
                continue
            cell = sheet.cell(row_index, col_index)
            target = worksheet.cell(row=row_index + 1, column=col_index + 1)
            target.value = _cell_value(cell, datemode)
            _apply_style(target, cell, xf_list, font_list, format_map)


def _cell_value(cell, datemode):
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_TEXT:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    if cell.ctype == xlrd.XL_CELL_BLANK:
        return None
    return cell.value


def _apply_style(target, cell, xf_list, font_list, format_map) -> None:
    xf_index = getattr(cell, "xf_index", None)
    if xf_index is None or xf_index >= len(xf_list):
        return
    xf = xf_list[xf_index]

    font = font_list[xf.font_index] if xf.font_index < len(font_list) else None
    if font is not None:
        color = _xls_color(font.colour_index)
        target.font = Font(
            name=font.name or "Calibri",
            size=font.height / 20.0 if font.height else 11,
            bold=bool(font.bold),
            italic=bool(font.italic),
            underline="single" if font.underline_type else None,
            strike=bool(font.struck_out),
            color=color,
        )

    align = xf.alignment
    horz_map = {
        1: "left",
        2: "center",
        3: "right",
        4: "fill",
        5: "justify",
        6: "centerContinuous",
        7: "distributed",
    }
    vert_map = {0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}
    
    # Manejo compatible para text_wrapped en lugar de wrap
    is_wrapped = bool(getattr(align, "text_wrapped", getattr(align, "wrap", 0)))

    target.alignment = Alignment(
        horizontal=horz_map.get(align.hor_align),
        vertical=vert_map.get(align.vert_align),
        wrap_text=is_wrapped,
        shrink_to_fit=bool(align.shrink_to_fit),
        text_rotation=align.rotation if 0 <= align.rotation <= 180 else 0,
    )

    fmt = format_map.get(xf.format_key)
    if fmt is not None and getattr(fmt, "format_str", None):
        excel_fmt = _translate_format(fmt.format_str)
        if excel_fmt:
            target.number_format = excel_fmt

    borders = xf.border
    side = lambda style, color_index: Side(
        style=_border_style(style),
        color=_xls_color(color_index) or "000000",
    ) if style else Side(style=None)
    target.border = Border(
        left=side(borders.left_line_style, borders.left_colour_index),
        right=side(borders.right_line_style, borders.right_colour_index),
        top=side(borders.top_line_style, borders.top_colour_index),
        bottom=side(borders.bottom_line_style, borders.bottom_colour_index),
    )

    bg = xf.background
    fill_color = _xls_color(bg.pattern_colour_index)
    if fill_color and bg.fill_pattern:
        target.fill = PatternFill(fill_type="solid", fgColor=fill_color)

    target.protection = Protection(locked=bool(xf.protection.cell_locked), hidden=bool(xf.protection.formula_hidden))


def _border_style(style_index: int) -> Optional[str]:
    mapping = {
        1: "thin",
        2: "medium",
        3: "dashed",
        4: "dotted",
        5: "thick",
        6: "double",
        7: "hair",
        8: "mediumDashed",
        9: "dashDot",
        10: "mediumDashDot",
        11: "dashDotDot",
        12: "mediumDashDotDot",
        13: "slantDashDot",
    }
    return mapping.get(style_index)


def _xls_color(index: Optional[int]) -> Optional[str]:
    palette = {
        0x08: "000000",
        0x09: "FFFFFF",
        0x0A: "FF0000",
        0x0B: "00FF00",
        0x0C: "0000FF",
        0x0D: "FFFF00",
        0x0E: "FF00FF",
        0x0F: "00FFFF",
        0x10: "800000",
        0x11: "008000",
        0x12: "000080",
        0x13: "808000",
        0x14: "800080",
        0x15: "008080",
        0x16: "C0C0C0",
        0x17: "808080",
        0x18: "9999FF",
        0x19: "993366",
        0x1A: "FFFFCC",
        0x1B: "CCFFFF",
        0x1C: "660066",
        0x1D: "FF8080",
        0x1E: "0066CC",
        0x1F: "CCCCFF",
        0x40: "000000",
        0x41: "FFFFFF",
    }
    if index is None:
        return None
    return palette.get(index)


def _translate_format(fmt: str) -> str:
    if not fmt:
        return "General"
    return fmt.replace("_-", "_").replace("\\", "")


def summarize(results: List[ConvertResult]) -> Tuple[int, int]:
    ok = sum(1 for item in results if item.ok)
    return ok, len(results) - ok