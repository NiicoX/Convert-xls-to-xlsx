# -*- coding: utf-8 -*-
from __future__ import unicode_literals

import tempfile
import unittest
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from xlwt import Workbook, easyxf

from converter import _convert_with_xlrd, convert_file, destination_for, discover_xls


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_converts_sheets_values_and_merged_cells(self):
        source = self._write_sample_xls(self.root / "reporte.xls")
        result = convert_file(source, overwrite=True)

        self.assertTrue(result.ok, result.message)
        self.assertTrue(result.destination.exists())

        book = load_workbook(str(result.destination))
        self.assertEqual(book.sheetnames, ["Datos", "Resumen"])

        sheet = book["Datos"]
        self.assertEqual(sheet["A1"].value, "Producto")
        self.assertEqual(sheet["B2"].value, 12)
        self.assertEqual(sheet["C2"].value, 19.5)
        self.assertTrue("A4:B4" in sheet.merged_cells)

    def test_xlrd_fallback_preserves_data(self):
        source = self._write_sample_xls(self.root / "sin_excel.xls")
        dest = self.root / "sin_excel.xlsx"
        _convert_with_xlrd(source, dest)
        book = load_workbook(str(dest))
        self.assertEqual(book["Datos"]["A2"].value, "Manzanas")
        self.assertEqual(book["Datos"]["B2"].value, 12)
        self.assertEqual(book.sheetnames, ["Datos", "Resumen"])

    def test_skips_existing_without_overwrite(self):
        source = self._write_sample_xls(self.root / "demo.xls")
        first = convert_file(source, overwrite=True)
        self.assertTrue(first.ok, first.message)
        second = convert_file(source, overwrite=False)
        self.assertFalse(second.ok)

    def test_discover_recursive(self):
        nested = self.root / "a" / "b"
        nested.mkdir(parents=True)
        self._write_sample_xls(self.root / "uno.xls")
        self._write_sample_xls(nested / "dos.xls")
        found = discover_xls(self.root, recursive=True)
        self.assertEqual(len(found), 2)
        self.assertEqual(destination_for(found[0]).suffix, ".xlsx")

    def _write_sample_xls(self, path):
        book = Workbook()
        header = easyxf("font: bold on; align: horiz center")
        currency = easyxf(num_format_str="0.00")

        datos = book.add_sheet("Datos")
        datos.write(0, 0, "Producto", header)
        datos.write(0, 1, "Cantidad", header)
        datos.write(0, 2, "Precio", header)
        datos.write(1, 0, "Manzanas")
        datos.write(1, 1, 12)
        datos.write(1, 2, 19.5, currency)
        datos.write(3, 0, "Total combinado")
        datos.merge(3, 3, 0, 1)

        resumen = book.add_sheet("Resumen")
        resumen.write(0, 0, "OK")

        book.save(str(path))
        xlrd.open_workbook(str(path))
        return path


if __name__ == "__main__":
    unittest.main()
